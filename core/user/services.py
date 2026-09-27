"""Lógica de negocio para la exportación e importación de permisos de grupos.

El archivo Excel generado contiene una hoja (pestaña) por cada grupo de
usuario existente. Cada hoja lista todos los permisos del sistema con una
columna "Asignado" que indica si el grupo posee dicho permiso, de modo que el
mismo archivo exportado puede editarse y volver a importarse (round-trip).
"""

import re

from django.contrib.auth.models import Group, Permission
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

GROUP_LABEL = 'Grupo de usuario'

PERMISSION_HEADERS = [
    'Aplicación',
    'Modelo',
    'Nombre del permiso',
    'Codename',
    'Asignado',
]

AFFIRMATIVE_VALUES = {'si', 'sí', 's', 'x', '1', 'true', 'verdadero', 'yes', 'y'}

INVALID_SHEET_CHARS = re.compile(r'[\[\]:*?/\\]')

MAX_SHEET_TITLE_LENGTH = 31


def _normalize(value):
    """Normaliza un valor de celda para comparaciones (minúsculas y sin espacios extra)."""
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip().lower()


def _unique_sheet_title(name, used_titles):
    """Genera un título de hoja válido y único para Excel a partir del nombre del grupo."""
    base = INVALID_SHEET_CHARS.sub('_', str(name or '')).strip().strip("'") or 'Grupo'
    base = base[:MAX_SHEET_TITLE_LENGTH]
    title = base
    counter = 1
    while title.lower() in used_titles:
        suffix = f'_{counter}'
        title = base[:MAX_SHEET_TITLE_LENGTH - len(suffix)] + suffix
        counter += 1
    used_titles.add(title.lower())
    return title


def build_group_permissions_workbook():
    """Construye un libro de Excel con una hoja por grupo y sus permisos.

    Cada hoja contiene en la fila 1 el nombre del grupo, en la fila 2 los
    encabezados y a partir de la fila 3 una fila por cada permiso del sistema
    indicando si está asignado al grupo. Retorna un objeto Workbook.
    """
    permissions = list(
        Permission.objects.select_related('content_type').order_by(
            'content_type__app_label', 'content_type__model', 'codename'
        )
    )
    groups = Group.objects.prefetch_related('permissions').order_by('name')

    workbook = Workbook()
    default_sheet = workbook.active

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', start_color='4472C4')
    group_font = Font(bold=True, color='1F3864')
    group_fill = PatternFill(fill_type='solid', start_color='D9E1F2')
    alignment = Alignment(vertical='center', wrap_text=False)

    used_titles = set()
    for group in groups:
        sheet = workbook.create_sheet(title=_unique_sheet_title(group.name, used_titles))

        title_cell = sheet.cell(row=1, column=1, value=GROUP_LABEL)
        title_cell.font = group_font
        title_cell.fill = group_fill
        name_cell = sheet.cell(row=1, column=2, value=group.name)
        name_cell.font = group_font
        name_cell.fill = group_fill

        for col_num, header in enumerate(PERMISSION_HEADERS, 1):
            cell = sheet.cell(row=2, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        assigned_ids = {permission.id for permission in group.permissions.all()}

        for row_num, permission in enumerate(permissions, 3):
            content_type = permission.content_type
            values = [
                content_type.app_label,
                content_type.model,
                permission.name,
                permission.codename,
                'Sí' if permission.id in assigned_ids else 'No',
            ]
            for col_num, value in enumerate(values, 1):
                cell = sheet.cell(row=row_num, column=col_num, value=value)
                cell.alignment = alignment

        sheet.freeze_panes = 'A3'
        widths = {'A': 24, 'B': 24, 'C': 45, 'D': 40, 'E': 12}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    if groups:
        workbook.remove(default_sheet)
    else:
        default_sheet.title = 'Sin grupos'
        default_sheet.cell(row=1, column=1, value='No hay grupos de usuario registrados.')

    return workbook


def _parse_sheet(sheet):
    """Extrae el nombre del grupo y los permisos de una hoja del archivo.

    Retorna un diccionario con el nombre del grupo y la lista de permisos
    encontrados, o None si la hoja no tiene el formato esperado.
    """
    group_name = None
    headers = None
    data_rows = []

    for row in sheet.iter_rows(values_only=True):
        if row is None or all(cell is None or str(cell).strip() == '' for cell in row):
            continue
        normalized = [_normalize(cell) for cell in row]

        if group_name is None and GROUP_LABEL.lower() in normalized:
            index = normalized.index(GROUP_LABEL.lower())
            if len(row) > index + 1 and row[index + 1] is not None and str(row[index + 1]).strip():
                group_name = str(row[index + 1]).strip()
                continue

        if headers is None and 'codename' in normalized and 'asignado' in normalized:
            headers = normalized
            continue

        if headers is not None:
            data_rows.append(row)

    if not group_name or headers is None:
        return None

    code_index = headers.index('codename')
    assigned_index = headers.index('asignado')
    app_index = headers.index('aplicación') if 'aplicación' in headers else None
    model_index = headers.index('modelo') if 'modelo' in headers else None
    required_index = max(code_index, assigned_index)

    permissions = []
    for row in data_rows:
        if len(row) <= required_index:
            continue
        codename = row[code_index]
        if codename is None or not str(codename).strip():
            continue
        app_label = ''
        model = ''
        if app_index is not None and len(row) > app_index and row[app_index] is not None:
            app_label = str(row[app_index]).strip()
        if model_index is not None and len(row) > model_index and row[model_index] is not None:
            model = str(row[model_index]).strip()
        permissions.append({
            'app_label': app_label,
            'model': model,
            'codename': str(codename).strip(),
            'assigned': _normalize(row[assigned_index]) in AFFIRMATIVE_VALUES,
        })

    return {'name': group_name, 'permissions': permissions}


def parse_group_permissions_workbook(excel_file):
    """Lee un archivo Excel y retorna la lista de grupos con sus permisos.

    Ignora las hojas que no tengan el formato esperado. Lanza ValueError si el
    archivo no es un Excel válido o no contiene ninguna hoja de grupos.
    """
    try:
        workbook = load_workbook(excel_file, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f'El archivo no es un Excel válido (.xlsx): {exc}')

    try:
        parsed_groups = []
        for sheet in workbook.worksheets:
            group_data = _parse_sheet(sheet)
            if group_data is not None:
                parsed_groups.append(group_data)
    finally:
        workbook.close()

    if not parsed_groups:
        raise ValueError(
            'El archivo no contiene hojas con el formato esperado. '
            'Descargue primero el archivo de permisos para usarlo como plantilla.'
        )
    return parsed_groups


def apply_group_permissions(parsed_groups):
    """Crea o actualiza los grupos y sus permisos a partir de los datos leídos.

    Para cada grupo se asignan los permisos marcados como afirmativos y se
    retiran los que aparezcan en la hoja marcados como negativos. Los permisos
    que no aparezcan en la hoja se dejan intactos. Retorna un resumen con los
    grupos creados, actualizados y las advertencias de permisos no encontrados.
    """
    codenames = {
        entry['codename']
        for group_data in parsed_groups
        for entry in group_data['permissions']
    }

    exact_lookup = {}
    by_codename = {}
    for permission in Permission.objects.select_related('content_type').filter(codename__in=codenames):
        content_type = permission.content_type
        exact_lookup[(content_type.app_label, content_type.model, permission.codename)] = permission
        by_codename.setdefault(permission.codename, []).append(permission)

    created = 0
    updated = 0
    warnings = []

    with transaction.atomic():
        for group_data in parsed_groups:
            group, group_created = Group.objects.get_or_create(name=group_data['name'])
            current_ids = set(group.permissions.values_list('id', flat=True))
            listed_ids = set()
            assign_ids = set()

            for entry in group_data['permissions']:
                permission = exact_lookup.get(
                    (entry['app_label'], entry['model'], entry['codename'])
                )
                if permission is None and not entry['app_label'] and not entry['model']:
                    matches = by_codename.get(entry['codename'], [])
                    permission = matches[0] if len(matches) == 1 else None
                if permission is None:
                    warnings.append(
                        f'Permiso no encontrado: {entry["app_label"]}.{entry["model"]} '
                        f'({entry["codename"]}).'
                    )
                    continue
                listed_ids.add(permission.id)
                if entry['assigned']:
                    assign_ids.add(permission.id)

            desired_ids = (current_ids - listed_ids) | assign_ids
            if desired_ids != current_ids:
                group.permissions.set(desired_ids)

            if group_created:
                created += 1
            elif desired_ids != current_ids:
                updated += 1

    return {'created': created, 'updated': updated, 'warnings': warnings}
